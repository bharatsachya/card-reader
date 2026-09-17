"""
End-to-end tests over the real ASGI app.

The model is the stub server from tools/, which speaks the same
/v1/chat/completions contract. That is the point of keeping model_client.py as
the only networked module: the whole pipeline -- upload, spool, preprocess,
prompt, parse, post-process, store, export -- is exercised for real, with only
the one non-deterministic step replaced.
"""

import asyncio
import io

import httpx
import pytest
from PIL import Image

from app.auth import User, require_user
from app.config import settings
from app.main import app
from app.store import store


def card_bytes(text: str = "Asha Rao") -> bytes:
    """A small in-memory JPEG. The stub ignores content, so pixels do not matter."""
    buffer = io.BytesIO()
    Image.new("RGB", (600, 360), "white").save(buffer, format="JPEG")
    return buffer.getvalue()


@pytest.fixture
async def client(stub_model, monkeypatch):
    """The app, wired to the stub model, with lifespan actually run."""
    monkeypatch.setattr(settings, "model_url", stub_model)
    monkeypatch.setattr(settings, "model_max_attempts", 1)   # fail fast in tests

    transport = httpx.ASGITransport(app=app)
    # Running the lifespan matters: it is what initialises the store and runs
    # the reclaim/sweep passes, so skipping it would test a different app.
    async with app.router.lifespan_context(app), httpx.AsyncClient(
        transport=transport, base_url="http://test"
    ) as http:
        yield http


async def run_job_to_completion(client, files, timeout=30.0):
    response = await client.post("/api/jobs", files=files)
    assert response.status_code == 202, response.text
    job_id = response.json()["job_id"]

    deadline = asyncio.get_running_loop().time() + timeout
    while True:
        poll = await client.get(f"/api/jobs/{job_id}?summary=true")
        assert poll.status_code == 200
        body = poll.json()
        if body["status"] in {"done", "failed"}:
            return job_id, body
        assert asyncio.get_running_loop().time() < deadline, "job never finished"
        await asyncio.sleep(0.05)


# --- the happy path, all the way through ----------------------------------

async def test_a_bulk_job_runs_end_to_end_and_exports(client):
    files = [("files", (f"card_{i}.jpg", card_bytes(), "image/jpeg")) for i in range(3)]
    job_id, summary = await run_job_to_completion(client, files)

    assert summary["status"] == "done"
    assert summary["total"] == 3 and summary["processed"] == 3

    full = (await client.get(f"/api/jobs/{job_id}")).json()
    assert len(full["leads"]) == 3
    # Every lead carries an id and a retained image, which is what the sidebar
    # thumbnails and the lightbox depend on.
    assert all(lead["id"] for lead in full["leads"])
    assert all(lead["image_sha256"] for lead in full["leads"])

    export = await client.get(f"/api/jobs/{job_id}/export.xlsx")
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert export.content[:2] == b"PK", "an xlsx is a zip archive"


async def test_a_non_image_becomes_a_flagged_row_not_a_500(client):
    files = [
        ("files", ("good.jpg", card_bytes(), "image/jpeg")),
        ("files", ("notes.txt", b"this is not an image", "text/plain")),
    ]
    job_id, summary = await run_job_to_completion(client, files)
    # The whole point: one bad item does not destroy the run.
    assert summary["status"] == "done" and summary["processed"] == 2
    leads = (await client.get(f"/api/jobs/{job_id}")).json()["leads"]
    statuses = {lead["source_filename"]: lead["status"] for lead in leads}
    assert statuses["notes.txt"] == "input_error"


async def test_too_many_files_is_refused_whole(client):
    files = [("files", (f"c{i}.jpg", card_bytes(), "image/jpeg"))
             for i in range(settings.max_files_per_request + 1)]
    response = await client.post("/api/jobs", files=files)
    assert response.status_code == 413


async def test_health_is_fast_and_reports_configuration(client):
    body = (await client.get("/health")).json()
    assert body["status"] == "ok"
    assert "model_endpoint" in body
    assert body["config"]["store_backend"] in {"memory", "sqlite"}


# --- image route ----------------------------------------------------------

async def test_the_retained_image_is_served_immutably(client):
    files = [("files", ("card.jpg", card_bytes(), "image/jpeg"))]
    job_id, _ = await run_job_to_completion(client, files)
    lead = (await client.get(f"/api/jobs/{job_id}")).json()["leads"][0]

    response = await client.get(f"/api/leads/{lead['id']}/image")
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    # private, not public: a shared cache must never hand one user's card to
    # another. immutable is only safe because the URL is content-addressed.
    cache = response.headers["cache-control"]
    assert "private" in cache and "immutable" in cache
    assert "public" not in cache
    assert response.content[:2] == b"\xff\xd8", "a real JPEG"


async def test_unknown_lead_image_is_404(client):
    assert (await client.get("/api/leads/deadbeefdeadbeef/image")).status_code == 404


# --- THE ISOLATION CASE, over HTTP ----------------------------------------

async def test_a_second_user_cannot_read_the_first_users_job_or_image(client):
    """
    Two users, one server. User B must not see User A's job, leads or images,
    and must not be able to tell a forbidden id from a nonexistent one.

    require_user is overridden rather than minting real Clerk tokens: this test
    is about what the ROUTES do with an identity, not about JWT verification,
    which auth.py covers separately. Overriding keeps the test hermetic and
    needs no Clerk account to run in CI.
    """
    def as_user(user_id):
        return lambda: User(id=user_id, email=f"{user_id}@example.com")

    app.dependency_overrides[require_user] = as_user("alice")
    try:
        files = [("files", ("card.jpg", card_bytes(), "image/jpeg"))]
        job_id, _ = await run_job_to_completion(client, files)
        lead = (await client.get(f"/api/jobs/{job_id}")).json()["leads"][0]

        alice_jobs = (await client.get("/api/jobs")).json()["jobs"]
        assert any(job["job_id"] == job_id for job in alice_jobs)

        # --- now become Bob ---
        app.dependency_overrides[require_user] = as_user("bob")

        # 404, NOT 403: a 403 would confirm the id exists, which is exactly
        # what lets an attacker enumerate valid ids.
        assert (await client.get(f"/api/jobs/{job_id}")).status_code == 404
        assert (await client.get(f"/api/jobs/{job_id}/export.xlsx")).status_code == 404
        assert (await client.get(f"/api/leads/{lead['id']}/image")).status_code == 404

        # Indistinguishable from an id that never existed.
        assert (await client.get("/api/jobs/000000000000")).status_code == 404

        # Bob's own views are empty, not merely filtered in the UI.
        assert (await client.get("/api/jobs")).json()["jobs"] == []
        bob_export = await client.get("/api/export.xlsx")
        assert bob_export.status_code == 200
        assert not await store.all_leads("bob")
    finally:
        app.dependency_overrides.clear()


async def test_missing_token_is_401_when_auth_is_enabled(client, monkeypatch):
    """With AUTH_MODE=clerk and no Authorization header, every data route 401s."""
    monkeypatch.setattr(settings, "auth_mode", "clerk")
    monkeypatch.setattr(settings, "clerk_issuer", "https://example.clerk.accounts.dev")
    monkeypatch.setattr(settings, "clerk_publishable_key", "pk_test_x")

    for path in ["/api/jobs", "/api/jobs/abc", "/api/leads/abc/image", "/api/export.xlsx"]:
        response = await client.get(path)
        assert response.status_code == 401, f"{path} returned {response.status_code}"
        assert response.json()["detail"]["code"] == "UNAUTHENTICATED"


async def test_a_forged_token_is_rejected(client, monkeypatch):
    """A self-signed HS256 token carries no Clerk signature, so it must fail."""
    import jwt as pyjwt

    monkeypatch.setattr(settings, "auth_mode", "clerk")
    monkeypatch.setattr(settings, "clerk_issuer", "https://example.clerk.accounts.dev")
    monkeypatch.setattr(settings, "clerk_publishable_key", "pk_test_x")

    forged = pyjwt.encode(
        {"sub": "attacker", "iss": "https://example.clerk.accounts.dev",
         "exp": 4102444800},
        "a-key-clerk-does-not-have-and-never-will-32b", algorithm="HS256",
    )
    response = await client.get(
        "/api/jobs", headers={"Authorization": f"Bearer {forged}"}
    )
    assert response.status_code == 401


# --- sessions -------------------------------------------------------------

async def test_two_uploads_in_one_session_make_one_spreadsheet(client):
    """
    The feature, stated as a test: upload twice into the same session and the
    session export contains everything, rather than two files to merge.
    """
    import openpyxl

    created = await client.post("/api/sessions")
    assert created.status_code == 201
    session_id = created.json()["session_id"]
    assert created.json()["title"], "a session needs a readable default title"

    first = [("files", ("a.jpg", card_bytes(), "image/jpeg"))]
    await run_job_to_completion(client, first + [("session_id", (None, session_id))])

    second = [("files", ("b.jpg", card_bytes(), "image/jpeg")),
              ("files", ("c.jpg", card_bytes(), "image/jpeg"))]
    await run_job_to_completion(client, second + [("session_id", (None, session_id))])

    detail = (await client.get(f"/api/sessions/{session_id}")).json()
    assert detail["jobs"] == 2, "both uploads belong to the session"
    assert detail["cards"] == 3

    export = await client.get(f"/api/sessions/{session_id}/export.xlsx")
    assert export.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
    # header + one row per card, from BOTH uploads
    assert sheet.max_row == 4, f"expected 3 data rows, got {sheet.max_row - 1}"


async def test_an_upload_without_a_session_opens_one(client):
    """A bare curl must still work -- no "create a session first" ceremony."""
    files = [("files", ("a.jpg", card_bytes(), "image/jpeg"))]
    response = await client.post("/api/jobs", files=files)
    assert response.status_code == 202
    assert response.json()["session_id"], "the server must say where it landed"


async def test_a_foreign_session_id_does_not_leak_or_crash(client):
    """
    Posting someone else's session id opens a fresh one rather than erroring.
    Rejecting it would confirm the id exists, which is the enumeration leak the
    404-not-403 rule elsewhere exists to prevent.
    """
    def as_user(uid):
        return lambda: User(id=uid, email=None)

    app.dependency_overrides[require_user] = as_user("alice")
    try:
        alice_session = (await client.post("/api/sessions")).json()["session_id"]

        app.dependency_overrides[require_user] = as_user("bob")
        files = [("files", ("a.jpg", card_bytes(), "image/jpeg"))]
        response = await client.post(
            "/api/jobs",
            files=files + [("session_id", (None, alice_session))],
        )
        assert response.status_code == 202
        assert response.json()["session_id"] != alice_session, "BOB JOINED ALICE'S SESSION"

        assert (await client.get(f"/api/sessions/{alice_session}")).status_code == 404
        assert (await client.get(
            f"/api/sessions/{alice_session}/export.xlsx")).status_code == 404
    finally:
        app.dependency_overrides.clear()


async def test_sessions_list_is_paginated_and_scoped(client):
    def as_user(uid):
        return lambda: User(id=uid, email=None)

    app.dependency_overrides[require_user] = as_user("carol")
    try:
        for _ in range(3):
            await client.post("/api/sessions")
        page = (await client.get("/api/sessions?limit=2")).json()
        assert len(page["sessions"]) == 2
        assert page["next_cursor"]
        rest = (await client.get(
            f"/api/sessions?limit=2&cursor={page['next_cursor']}")).json()
        assert len(rest["sessions"]) == 1
        assert rest["next_cursor"] is None
    finally:
        app.dependency_overrides.clear()
